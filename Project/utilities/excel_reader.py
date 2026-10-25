from openpyxl import load_workbook
from utilities.config_reader import ConfigReader
class ExcelReader:
 
    def __init__(self):
        """
        Method name: __init__
        Author name: Grace
        Short description of method:
            Loads the Excel workbook from the configured path.
        Return type: None
        Parameter list: None
        """
        try:
            config=ConfigReader()
            path=config.get_value("PATH","excel_path")
            self.workbook = load_workbook(path)
        except Exception as e:
            raise Exception(f"Failed to load Excel workbook: {e}")
 
    def get_data(self, sheet_name, row, col):
        """
        Method name: get_data
        Author name: Grace
        Short description of method:
            Fetches data from a given Excel sheet, row, and column.
        Return type: Any
        Parameter list:
            sheet_name (str)
            row (int)
            col (int)
        """
        try:
            sheet = self.workbook[sheet_name]
            data = sheet.cell(row=row, column=col).value
            return data
        except Exception as e:
            raise Exception(
                f"Failed to read data from sheet '{sheet_name}', "
                f"row {row}, column {col}: {e}"
            )